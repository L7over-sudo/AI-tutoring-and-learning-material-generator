from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "AI自动作图工具使用说明书.docx"
OUTPUT = ROOT / "AI自动作图工具图文使用说明书.docx"
FIG = ROOT / "figures"


SECTIONS = [
    ("1 软件启动与激活", "00-activation-client.png", "客户端激活页：粘贴激活码后立即激活"),
    ("2 主界面说明", "01-main.png", "主界面功能分区与常用入口"),
    ("3 任务列表与生成流程", "01-main.png", "任务列表、参数区、预览区和底部操作栏"),
    ("4 导入图片", "02-import-images.png", "导入图片或整个素材文件夹"),
    ("5 图文生图", "03-text-image.png", "填写主题并上传参考图生成新任务"),
    ("6 文档生图", "04-doc-images.png", "上传文档并转换为逐页图片"),
    ("7 文档生PPT", "05-doc-ppt.png", "从文档规划、生成并导出 PPT"),
    ("8 格式转换", "06-convert.png", "文档、图片、PDF 与 PPT 互转"),
    ("9 Word拆分", "07-word-split.png", "按目标字符数拆分 Word 文档"),
    ("10 PDF排序", "08-pdf-sort.png", "重排、倒序或随机排列 PDF 页面"),
    ("11 PDF水印", "09-pdf-watermark.png", "为 PDF 添加文字或图片水印"),
    ("12 图片水印", "10-image-watermark.png", "批量添加图片文字水印"),
    ("13 Markdown编辑器", "11-markdown.png", "Markdown 编辑、分栏预览与导出"),
    ("14 HTML编辑器", "12-html.png", "HTML、CSS、JavaScript 编辑与实时预览"),
    ("15 分镜", "13-storyboard.png", "导入文案、拆分分镜并加入生成任务"),
    ("16 图片拼接", "14-stitch.png", "图片排序、版式设置与拼接下载"),
    ("17 作品解析", "15-xhs-parse.png", "解析公开作品链接并保存素材"),
    ("18 提示词模板", "17-prompt-templates.png", "提示词模板的新建、编辑和应用"),
    ("19 发布小红书", "01-main.png", "底部操作栏中的“发布小红书”入口"),
    ("20 商品上架", "01-main.png", "底部操作栏中的“上架商品”入口"),
    ("21 API设置", "16-api-settings.png", "填写接口参数、测试连接并保存"),
    ("22 激活码发卡后台", "18-license-admin.png", "选择卡型，一键生成并复制激活码"),
    ("管理发卡记录", "19-activation-records.png", "搜索、筛选并查看激活码状态"),
]


def insert_figure_after(doc: Document, anchor, image_path: Path, caption: str, number: int):
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(3)
    picture.paragraph_format.space_after = Pt(2)
    picture.paragraph_format.keep_with_next = True
    run = picture.add_run()
    run.add_picture(str(image_path), width=Inches(6.25))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_with_next = True
    cr = cap.add_run(f"操作图 {number}｜{caption}")
    cr.font.name = "Microsoft YaHei"
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(85, 108, 102)

    anchor._p.addnext(picture._p)
    picture._p.addnext(cap._p)


doc = Document(SOURCE)

# Refresh the cover so the imported Feishu version is clearly distinguishable.
for p in doc.paragraphs[:10]:
    if p.text.strip() == "AI自动作图工具使用说明书":
        p.text = "AI自动作图工具图文使用说明书"
    elif p.text.strip().startswith("软件版本"):
        p.text = "软件版本 1.0.7｜图文版 1.1"

# Start each major feature on a clean page. This keeps screenshots and steps together.
for p in doc.paragraphs:
    text = p.text.strip()
    if p.style.name == "Heading 1" and text[:1].isdigit():
        p.paragraph_format.page_break_before = True
    if p.style.name in {"Heading 1", "Heading 2"}:
        p.paragraph_format.keep_with_next = True

anchors = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}
for idx, (anchor_text, filename, caption) in enumerate(SECTIONS, 1):
    anchor = anchors.get(anchor_text)
    if anchor is None:
        raise RuntimeError(f"Missing section anchor: {anchor_text}")
    image_path = FIG / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    insert_figure_after(doc, anchor, image_path, caption, idx)

# Add a short visual legend below the quick-start heading.
quick = anchors.get("一分钟快速上手")
if quick is not None:
    legend = doc.add_paragraph()
    legend.paragraph_format.space_after = Pt(6)
    lr = legend.add_run("图中红色编号表示推荐操作顺序；截图中的密钥均已遮挡，实际使用时请填写自己的账号配置。")
    lr.font.name = "Microsoft YaHei"
    lr.font.size = Pt(9.5)
    lr.font.color.rgb = RGBColor(70, 95, 88)
    quick._p.addnext(legend._p)

doc.core_properties.title = "AI自动作图工具图文使用说明书"
doc.core_properties.subject = "AI自动作图工具全功能图文操作指南"
doc.core_properties.comments = "图文版 1.1，截图来自本地客户端实测页面。"
doc.save(OUTPUT)
print(OUTPUT)
